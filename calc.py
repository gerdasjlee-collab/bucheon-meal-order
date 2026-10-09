# -*- coding: utf-8 -*-
"""식단 → 인원 → 구매량 계산 엔진 + 엑셀 출력."""
import io
import math

import pandas as pd
import streamlit as st

from waste import default_waste

MEAL_ORDER = ["조식", "오전간식", "중식", "오후간식", "석식"]
WEEKDAY = ["월", "화", "수", "목", "금", "토", "일"]


@st.cache_data(show_spinner=False)
def load_db():
    d = pd.read_csv("db_dishes.csv.gz", dtype={"조리방법": str}).fillna({"알레르기": "", "조리방법": ""})
    i = pd.read_csv("db_ingredients.csv.gz")
    d["일"] = d["일"].astype(int)
    return d, i


def weekday_of(month: str, day: int) -> str:
    import datetime as dt
    y, m = map(int, month.split("-"))
    try:
        return WEEKDAY[dt.date(y, m, day).weekday()]
    except ValueError:
        return ""


def days_in_week(dishes: pd.DataFrame, menu_type: str, week: str) -> list:
    return sorted(dishes[(dishes["식단유형"] == menu_type) & (dishes["주차"] == week)]["일"].unique().tolist())


def select_dishes(dishes: pd.DataFrame, menu_type: str, days: list, meals: list) -> pd.DataFrame:
    sel = dishes[(dishes["식단유형"] == menu_type) & (dishes["일"].isin(days)) & (dishes["끼니"].isin(meals))]
    return sel.sort_values(["일", "끼니"], key=lambda s: s.map(MEAL_ORDER.index) if s.name == "끼니" else s)


def compute(dishes, ings, portions: list, days: list, meals: list, extra: list = None):
    """portions: [(식단유형, 환산인원, 라벨)], extra: [(식단유형, 일, 끼니, 음식명, [(식재료,g)], 환산인원)] (AI/직접 추가 메뉴)
    반환: (식재료 집계 DataFrame, 음식별 상세 DataFrame)"""
    rows = []
    for menu_type, persons, label in portions:
        if persons <= 0:
            continue
        sel = select_dishes(dishes, menu_type, days, meals)
        m = sel.merge(ings, on="dish_id")
        m["환산인원"] = persons
        m["대상"] = label
        rows.append(m)
    for menu_type, day, meal, dish, ing_list, persons in (extra or []):
        for ing, g in ing_list:
            rows.append(pd.DataFrame([{"dish_id": -1, "식단유형": menu_type, "일": day, "끼니": meal, "음식명": dish + " (추가)", "알레르기": "", "조리방법": "", "식재료": ing, "g": g, "환산인원": persons, "대상": label_for(menu_type)}]))
    if not rows:
        return pd.DataFrame(), pd.DataFrame()
    m = pd.concat(rows, ignore_index=True)
    m["정미량g"] = m["g"] * m["환산인원"]
    detail = m[["일", "끼니", "음식명", "알레르기", "대상", "식재료", "g", "환산인원", "정미량g"]].copy()
    agg = m.groupby("식재료", as_index=False).agg(정미량g=("정미량g", "sum"), 사용음식수=("음식명", "nunique"), 사용음식=("음식명", lambda s: ", ".join(sorted(set(s))[:6])))
    agg["폐기율%"] = agg["식재료"].map(default_waste)
    agg = agg.sort_values("정미량g", ascending=False).reset_index(drop=True)
    return agg, detail


def label_for(menu_type: str) -> str:
    return menu_type


def apply_purchase(agg: pd.DataFrame, margin_pct: float) -> pd.DataFrame:
    """폐기율·여유율·단가 반영 → 구매량(kg)·금액."""
    a = agg.copy()
    a["폐기율%"] = pd.to_numeric(a["폐기율%"], errors="coerce").fillna(0).clip(0, 95)
    if "단가(원/kg)" not in a.columns:
        a["단가(원/kg)"] = 0.0
    a["단가(원/kg)"] = pd.to_numeric(a["단가(원/kg)"], errors="coerce").fillna(0)
    a["구매량g"] = a["정미량g"] * (1 + margin_pct / 100) / (1 - a["폐기율%"] / 100)
    a["구매량(kg)"] = (a["구매량g"] / 1000).apply(lambda x: math.ceil(x * 100) / 100)  # 10 g 단위 올림
    a["정미량(kg)"] = (a["정미량g"] / 1000).round(3)
    a["금액(원)"] = (a["구매량(kg)"] * a["단가(원/kg)"]).round(0)
    return a


def recipe_cards(detail: pd.DataFrame, dishes: pd.DataFrame) -> list:
    """음식별 카드: 재료 1인분·환산량 + 조리방법."""
    cards = []
    if detail.empty:
        return cards
    how = dishes.set_index(["식단유형", "일", "끼니", "음식명"])["조리방법"].to_dict() if "식단유형" in dishes else {}
    for (day, meal, dish, who), g in detail.groupby(["일", "끼니", "음식명", "대상"], sort=False):
        persons = int(g["환산인원"].iloc[0])
        cards.append(dict(일=day, 끼니=meal, 음식명=dish, 대상=who, 환산인원=persons, 알레르기=g["알레르기"].iloc[0],
                          재료=[(r["식재료"], r["g"], round(r["정미량g"], 1)) for _, r in g.iterrows()],
                          조리방법=how.get((who, day, meal, dish), "")))
    return cards


def menu_table(dishes: pd.DataFrame, menu_types: list, days: list, meals: list, month: str) -> pd.DataFrame:
    """식단표: 행=끼니, 열=일자. 여러 유형이면 유형명 접두."""
    sel = dishes[(dishes["식단유형"].isin(menu_types)) & (dishes["일"].isin(days)) & (dishes["끼니"].isin(meals))]
    cols = [f"{d}일({weekday_of(month, d)})" for d in days]
    out = {}
    for meal in [m for m in MEAL_ORDER if m in meals]:
        row = []
        for d in days:
            s = sel[(sel["일"] == d) & (sel["끼니"] == meal)]
            names, seen = [], {}
            for _, r in s.iterrows():
                tag = r["음식명"] + (r["알레르기"] if r["알레르기"] else "")
                seen.setdefault(tag, set()).add(r["식단유형"])
            for tag, types_ in seen.items():
                if len(menu_types) > 1 and len(types_) < len(menu_types):
                    tag = f"[{'/'.join(short(t) for t in sorted(types_))}] " + tag
                names.append(tag)
            row.append("\n".join(names))
        out[meal] = row
    return pd.DataFrame(out, index=cols).T


def short(menu_type: str) -> str:
    return menu_type.replace("어린이집 ", "").replace("(노인복지시설)", "").replace("(재가노인복지시설)", "재가").replace("아동식 ", "")[:14]


def to_excel(info: dict, purchase: pd.DataFrame, detail: pd.DataFrame, cards: list, menu: pd.DataFrame) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8F0FE")

    def header(ws, row, vals):
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=j, value=v)
            c.font = bold
            c.fill = fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 1) 발주서
    ws = wb.active
    ws.title = "발주서"
    ws["A1"] = f"식재료 발주서 — {info.get('기관명', '')}"
    ws["A1"].font = Font(bold=True, size=14)
    r = 2
    for k, v in info.items():
        ws.cell(row=r, column=1, value=k).font = bold
        ws.cell(row=r, column=2, value=str(v))
        r += 1
    r += 1
    cols = ["식재료", "정미량(kg)", "폐기율%", "구매량(kg)", "단가(원/kg)", "금액(원)"] + (["단가출처"] if "단가출처" in purchase.columns else []) + ["사용음식"]
    header(ws, r, cols)
    for _, p in purchase.iterrows():
        r += 1
        for j, c in enumerate(cols, 1):
            ws.cell(row=r, column=j, value=p[c])
    r += 1
    ws.cell(row=r, column=1, value="합계").font = bold
    ws.cell(row=r, column=4, value=round(float(purchase["구매량(kg)"].sum()), 2)).font = bold
    ws.cell(row=r, column=6, value=float(purchase["금액(원)"].sum())).font = bold
    for j, w in enumerate([30, 12, 10, 12, 12, 14, 50], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=len(info) + 4, column=1)

    # 2) 식단표
    ws2 = wb.create_sheet("식단표")
    ws2["A1"] = f"식단표 — {info.get('식단', '')}"
    ws2["A1"].font = Font(bold=True, size=13)
    header(ws2, 3, ["끼니"] + list(menu.columns))
    for i, (meal, row) in enumerate(menu.iterrows(), 4):
        ws2.cell(row=i, column=1, value=meal).font = bold
        for j, v in enumerate(row.tolist(), 2):
            c = ws2.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.row_dimensions[i].height = 110
    for j in range(2, len(menu.columns) + 2):
        ws2.column_dimensions[get_column_letter(j)].width = 24
    ws2.column_dimensions["A"].width = 10

    # 3) 레시피 카드
    ws3 = wb.create_sheet("레시피카드")
    header(ws3, 1, ["일", "끼니", "음식명", "알레르기", "대상", "환산인원", "식재료", "1인분(g)", "총 정미량(g)", "조리방법"])
    r = 1
    for cd in cards:
        for k, (ing, g1, gt) in enumerate(cd["재료"]):
            r += 1
            vals = [cd["일"], cd["끼니"], cd["음식명"], cd["알레르기"], cd["대상"], cd["환산인원"], ing, g1, gt, cd["조리방법"] if k == 0 else ""]
            for j, v in enumerate(vals, 1):
                c = ws3.cell(row=r, column=j, value=v)
                if j == 10:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate([6, 9, 22, 10, 26, 9, 28, 10, 12, 70], 1):
        ws3.column_dimensions[get_column_letter(j)].width = w
    ws3.freeze_panes = "A2"

    # 4) 음식별 상세
    ws4 = wb.create_sheet("음식별재료")
    header(ws4, 1, list(detail.columns))
    for i, row in enumerate(detail.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            ws4.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
